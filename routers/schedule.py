"""
Schedule Router - Routes for the Schedule feature
User-configurable tables for tracking work outside main project scope
"""
from fastapi import APIRouter, Request, HTTPException, UploadFile, File, Form
from fastapi.responses import HTMLResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from pathlib import Path
from typing import Dict, Any, List, Optional
from pydantic import BaseModel
import logging
import os
import io
import json
import csv
import re
import yaml
from datetime import date as _date, datetime as _datetime

from repositories.schedule_repository import ScheduleRepository

logger = logging.getLogger(__name__)

router = APIRouter()

# Setup templates
BASE_DIR = Path(__file__).resolve().parent.parent
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))

# Initialize repository
DATA_STORAGE_PATH = os.getenv("DATA_STORAGE_PATH", str(BASE_DIR / "mock_data"))
schedule_repo = ScheduleRepository(Path(DATA_STORAGE_PATH))


def get_build_version():
    """Get build version without circular import"""
    try:
        import main
        return main.BUILD_VERSION
    except:
        return "unknown"


# Pydantic models for request validation
class ColumnConfig(BaseModel):
    id: Optional[str] = None
    header: str
    type: str  # text, dropdown, date, number
    options: Optional[List[str]] = None  # for dropdown type
    width: Optional[int] = 150
    visible_in_export: bool = True


class CreateTableRequest(BaseModel):
    name: str
    description: Optional[str] = None
    columns: Optional[List[ColumnConfig]] = None
    color: Optional[str] = None


class UpdateTableRequest(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    columns: Optional[List[Dict[str, Any]]] = None
    rows: Optional[List[Dict[str, Any]]] = None
    color: Optional[str] = None


class RowData(BaseModel):
    data: Dict[str, Any]


class CopyTableRequest(BaseModel):
    source_project: str
    source_table_id: str
    new_table_name: Optional[str] = None
    include_data: bool = True


def get_user_from_request(request: Request):
    """Get user from request state"""
    return getattr(request.state, 'user', None)


@router.get("/schedule", response_class=HTMLResponse)
async def schedule_page(request: Request, project: str = None):
    """
    Main schedule page with sub-tabs for different schedule tables
    """
    if not project:
        # Redirect to project selection
        return templates.TemplateResponse("select_project.html", {
            "request": request,
            "redirect_to": "/dashboard/schedule",
            "build_version": get_build_version(),
            "user": get_user_from_request(request)
        })
    
    # Clean project name
    import re
    clean_name = project.replace('.xml', '').replace('.xlsx', '').replace('.yaml', '').strip()
    clean_name = re.sub(r'-\d+$', '', clean_name).strip()
    
    # Get schedule data
    schedule_data = schedule_repo.get_schedules(clean_name)
    
    # CSRF token from middleware for DELETE/POST/PUT requests
    csrf_token = getattr(request.state, 'csrf_token', '')
    
    context = {
        "request": request,
        "project_name": project,
        "clean_name": clean_name,
        "schedule_data": schedule_data,
        "tables": schedule_data.get('tables', []),
        "build_version": get_build_version(),
        "user": get_user_from_request(request),
        "csrf_token": csrf_token
    }
    
    response = templates.TemplateResponse("schedule.html", context)
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    return response


# =============================================================================
# Cross-Program Table Operations (must be before {project_name} routes)
# =============================================================================

@router.get("/api/schedule/all-programs/tables")
async def get_all_programs_with_tables():
    """
    Get all programs that have schedule tables.
    Used for "Copy from another program" feature.
    
    Returns:
        List of programs with their tables (name and id only, no row data)
    """
    try:
        programs_with_tables = []
        
        # Scan the schedules directory for all schedule files
        if schedule_repo.storage_dir.exists():
            for schedule_file in schedule_repo.storage_dir.glob("*_schedules.yaml"):
                try:
                    with open(schedule_file, 'r', encoding='utf-8') as f:
                        data = yaml.safe_load(f) or {}
                    
                    tables = data.get('tables', [])
                    if tables:  # Only include programs that have tables
                        project_name = data.get('project_name', schedule_file.stem.replace('_schedules', ''))
                        programs_with_tables.append({
                            'project_name': project_name,
                            'tables': [{'id': t['id'], 'name': t['name'], 'row_count': len(t.get('rows', []))} 
                                      for t in tables]
                        })
                except Exception as e:
                    logger.warning(f"Error reading schedule file {schedule_file}: {e}")
                    continue
        
        return JSONResponse(content={"programs": programs_with_tables})
        
    except Exception as e:
        logger.error(f"Error getting all programs with tables: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# =============================================================================
# API Endpoints for Schedule Tables
# =============================================================================

@router.get("/api/schedule/{project_name}")
async def get_project_schedules(project_name: str):
    """Get all schedule tables for a project"""
    logger.info(f"📋 API /api/schedule/{project_name} called")
    data = schedule_repo.get_schedules(project_name)
    tables = data.get('tables', [])
    logger.info(f"📋 Returning {len(tables)} tables for project '{project_name}'")
    return JSONResponse(content=data)


@router.post("/api/schedule/{project_name}/tables")
async def create_schedule_table(project_name: str, request: CreateTableRequest):
    """Create a new schedule table"""
    columns = None
    if request.columns:
        columns = [col.dict() for col in request.columns]
    
    table = schedule_repo.create_table(project_name, request.name, columns, description=request.description, color=request.color)
    return JSONResponse(content={"success": True, "table": table})


@router.get("/api/schedule/{project_name}/tables/{table_id}")
async def get_schedule_table(project_name: str, table_id: str):
    """Get a specific schedule table"""
    table = schedule_repo.get_table(project_name, table_id)
    if not table:
        raise HTTPException(status_code=404, detail="Table not found")
    return JSONResponse(content=table)


@router.put("/api/schedule/{project_name}/tables/{table_id}")
async def update_schedule_table(project_name: str, table_id: str, request: UpdateTableRequest):
    """Update a schedule table (name, columns, or rows)"""
    updates = request.dict(exclude_none=True)
    success = schedule_repo.update_table(project_name, table_id, updates)
    if not success:
        raise HTTPException(status_code=404, detail="Table not found")
    return JSONResponse(content={"success": True})


@router.delete("/api/schedule/{project_name}/tables/{table_id}")
async def delete_schedule_table(project_name: str, table_id: str):
    """Delete a schedule table"""
    try:
        logger.info(f"Deleting table {table_id} from project {project_name}")
        success = schedule_repo.delete_table(project_name, table_id)
        
        if not success:
            raise HTTPException(status_code=404, detail="Table not found")
        
        return JSONResponse(content={"success": True})
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error deleting table: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# =============================================================================
# Row Operations
# =============================================================================

@router.post("/api/schedule/{project_name}/tables/{table_id}/rows")
async def add_table_row(project_name: str, table_id: str, request: RowData = None):
    """Add a new row to a schedule table"""
    row_data = request.data if request else None
    row = schedule_repo.add_row(project_name, table_id, row_data)
    if not row:
        raise HTTPException(status_code=404, detail="Table not found")
    return JSONResponse(content={"success": True, "row": row})


@router.put("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}")
async def update_table_row(project_name: str, table_id: str, row_id: str, request: RowData):
    """Update a row in a schedule table"""
    success = schedule_repo.update_row(project_name, table_id, row_id, request.data)
    if not success:
        raise HTTPException(status_code=404, detail="Row not found")
    return JSONResponse(content={"success": True})


@router.patch("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/reschedule")
async def reschedule_row(project_name: str, table_id: str, row_id: str, request: Request):
    """Update the date of a single schedule row (called from calendar drag-and-drop)"""
    body = await request.json()
    date_col_id = body.get('date_col_id')
    new_date = body.get('new_date')

    if not date_col_id or not new_date:
        raise HTTPException(status_code=400, detail="date_col_id and new_date required")

    table = schedule_repo.get_table(project_name, table_id)
    if not table:
        raise HTTPException(status_code=404, detail="Table not found")

    row = next((r for r in table.get('rows', []) if r.get('id') == row_id), None)
    if not row:
        raise HTTPException(status_code=404, detail="Row not found")

    row_data = dict(row.get('data', {}))
    row_data[date_col_id] = new_date

    success = schedule_repo.update_row(project_name, table_id, row_id, row_data)
    try:
        from routers.calendar import invalidate_calendar_cache
        invalidate_calendar_cache()
    except Exception:
        pass
    return JSONResponse(content={"success": success})


@router.patch("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/cells")
async def update_row_cells(project_name: str, table_id: str, row_id: str, request: Request):
    """Update one or more cells in a schedule row by column ID."""
    body = await request.json()
    updates = body.get('updates')  # dict of {col_id: new_value}
    if not updates or not isinstance(updates, dict):
        raise HTTPException(status_code=400, detail="'updates' dict of {col_id: value} required")

    table = schedule_repo.get_table(project_name, table_id)
    if not table:
        raise HTTPException(status_code=404, detail="Table not found")

    row = next((r for r in table.get('rows', []) if r.get('id') == row_id), None)
    if not row:
        raise HTTPException(status_code=404, detail="Row not found")

    # Validate col_ids against table columns
    valid_col_ids = {c.get('id') for c in table.get('columns', [])}
    for col_id in updates:
        if col_id not in valid_col_ids:
            raise HTTPException(status_code=400, detail=f"Unknown column ID: {col_id}")

    row_data = dict(row.get('data', {}))
    for col_id, value in updates.items():
        row_data[col_id] = str(value) if value is not None else ''

    success = schedule_repo.update_row(project_name, table_id, row_id, row_data)
    try:
        from routers.calendar import invalidate_calendar_cache
        invalidate_calendar_cache()
    except Exception:
        pass
    return JSONResponse(content={"success": success})


@router.patch("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/complete")
async def complete_schedule_row(project_name: str, table_id: str, row_id: str):
    """Mark a schedule row as complete (used from calendar 'Done' button)"""
    table = schedule_repo.get_table(project_name, table_id)
    if not table:
        raise HTTPException(status_code=404, detail="Table not found")

    row = next((r for r in table.get('rows', []) if r.get('id') == row_id), None)
    if not row:
        raise HTTPException(status_code=404, detail="Row not found")

    row_data = dict(row.get('data', {}))

    # Find the status column (type == status/dropdown, or header contains 'status')
    columns = table.get('columns', [])
    status_col_id = None
    for col in columns:
        col_type = col.get('type', '')
        col_header = (col.get('header', '') or '').lower()
        if col_type in ('dropdown', 'status') or 'status' in col_header:
            status_col_id = col.get('id')
            break

    if not status_col_id:
        raise HTTPException(status_code=400, detail="No status column found in table")

    row_data[status_col_id] = 'Complete'
    success = schedule_repo.update_row(project_name, table_id, row_id, row_data)
    try:
        from routers.calendar import invalidate_calendar_cache
        invalidate_calendar_cache()
    except Exception:
        pass
    return JSONResponse(content={"success": success})


@router.patch("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/notes")
async def update_schedule_row_notes(project_name: str, table_id: str, row_id: str, request: Request):
    """Update the notes on a schedule row (stored separately from column data)."""
    body = await request.json()
    notes = body.get('notes', '')
    success = schedule_repo.update_row_notes(project_name, table_id, row_id, notes)
    if not success:
        raise HTTPException(status_code=404, detail="Row not found")
    return JSONResponse(content={"success": True})


@router.post("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/sub-tasks")
async def add_schedule_sub_task(project_name: str, table_id: str, row_id: str, request: Request):
    """Add a sub-task to a schedule row."""
    body = await request.json()
    title = (body.get('title') or '').strip()
    if not title:
        raise HTTPException(status_code=400, detail="Sub-task title is required")
    notes = (body.get('notes') or '').strip()
    sub_task = schedule_repo.add_sub_task(project_name, table_id, row_id, title, notes=notes)
    if not sub_task:
        raise HTTPException(status_code=404, detail="Row not found")
    from routers.calendar import invalidate_calendar_cache
    invalidate_calendar_cache()
    return JSONResponse(content={"success": True, "sub_task": sub_task})


@router.patch("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/sub-tasks/{sub_task_id}")
async def update_schedule_sub_task(project_name: str, table_id: str, row_id: str,
                                    sub_task_id: str, request: Request):
    """Update a sub-task (toggle completed, rename, notes)."""
    body = await request.json()
    updates = {}
    if 'completed' in body:
        updates['completed'] = bool(body['completed'])
    if 'title' in body:
        updates['title'] = str(body['title']).strip()
    if 'notes' in body:
        updates['notes'] = str(body['notes']).strip()
    if not updates:
        raise HTTPException(status_code=400, detail="No updates provided")
    success = schedule_repo.update_sub_task(project_name, table_id, row_id, sub_task_id, updates)
    if not success:
        raise HTTPException(status_code=404, detail="Sub-task not found")
    from routers.calendar import invalidate_calendar_cache
    invalidate_calendar_cache()
    return JSONResponse(content={"success": True})


@router.delete("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/sub-tasks/{sub_task_id}")
async def delete_schedule_sub_task(project_name: str, table_id: str, row_id: str, sub_task_id: str):
    """Delete a sub-task from a schedule row."""
    success = schedule_repo.delete_sub_task(project_name, table_id, row_id, sub_task_id)
    if not success:
        raise HTTPException(status_code=404, detail="Sub-task not found")
    from routers.calendar import invalidate_calendar_cache
    invalidate_calendar_cache()
    return JSONResponse(content={"success": True})


@router.put("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}/sub-tasks/reorder")
async def reorder_schedule_sub_tasks(project_name: str, table_id: str, row_id: str, request: Request):
    """Reorder sub-tasks on a schedule row. Body: { "order": ["id1","id2",...] }"""
    body = await request.json()
    ordered_ids = body.get('order')
    if not isinstance(ordered_ids, list):
        raise HTTPException(status_code=400, detail="'order' must be a list of sub-task IDs")
    success = schedule_repo.reorder_sub_tasks(project_name, table_id, row_id, ordered_ids)
    if not success:
        raise HTTPException(status_code=404, detail="Row not found")
    return JSONResponse(content={"success": True})


@router.delete("/api/schedule/{project_name}/tables/{table_id}/rows/{row_id}")
async def delete_table_row(project_name: str, table_id: str, row_id: str):
    """Delete a row from a schedule table"""
    success = schedule_repo.delete_row(project_name, table_id, row_id)
    if not success:
        raise HTTPException(status_code=404, detail="Row not found")
    return JSONResponse(content={"success": True})


# =============================================================================
# Export Views
# =============================================================================

@router.get("/schedule/table/{project_name}/{table_id}", response_class=HTMLResponse)
async def schedule_table_view(
    project_name: str, 
    table_id: str,
    export: bool = False,
    ppt_export: bool = False
):
    """
    Render a schedule table for viewing/export
    
    Args:
        project_name: Project name
        table_id: Table ID to render
        export: If True, only show columns marked visible_in_export
        ppt_export: If True, format for PowerPoint slide
    """
    table = schedule_repo.get_table(project_name, table_id)
    if not table:
        return HTMLResponse(
            content=f"<html><body><h2>Table not found</h2></body></html>",
            status_code=404
        )
    
    # Filter columns for export if needed
    columns = table.get('columns', [])
    if export or ppt_export:
        columns = [c for c in columns if c.get('visible_in_export', True)]
    
    # Build HTML table
    # For print/PDF view (not slide), add print controls and auto-trigger print dialog
    is_print_view = export and not ppt_export
    
    html = f'''<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{table.get('name', 'Schedule')} - PDF Export</title>
    <style>
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        {'html { height: 100%; }' if ppt_export else ''}
        body {{ 
            font-family: Arial, sans-serif; 
            background: white; 
            padding: {'40px 60px' if ppt_export else '20px 30px'};
            {'height: 100vh;' if ppt_export else 'min-height: 100vh;'}
            {'display: flex; flex-direction: column;' if ppt_export else ''}
        }}
        .slide-title {{
            color: #7F7F7F;
            font-size: 32px;
            margin-bottom: 24px;
            font-weight: normal;
            {'display: none;' if ppt_export else ''}
        }}
        .print-toolbar {{
            display: flex;
            align-items: center;
            gap: 12px;
            margin-bottom: 20px;
            padding: 12px 16px;
            background: #F3F4F6;
            border-radius: 8px;
            border: 1px solid #E5E7EB;
        }}
        .print-toolbar h2 {{
            flex: 1;
            font-size: 18px;
            color: #374151;
            margin: 0;
        }}
        .print-toolbar button {{
            padding: 8px 18px;
            border: none;
            border-radius: 6px;
            cursor: pointer;
            font-size: 14px;
            font-weight: 500;
            transition: background 0.15s;
        }}
        .btn-print {{
            background: #1E40AF;
            color: white;
        }}
        .btn-print:hover {{ background: #1E3A8A; }}
        .btn-close {{
            background: #E5E7EB;
            color: #374151;
        }}
        .btn-close:hover {{ background: #D1D5DB; }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: {'20px' if ppt_export else '16px'};
            table-layout: fixed;
            {'flex: 1;' if ppt_export else ''}
        }}
        th {{
            background: #1E40AF;
            color: white;
            padding: {'16px 14px' if ppt_export else '12px 10px'};
            text-align: left;
            font-weight: bold;
            font-size: {'19px' if ppt_export else '15px'};
            white-space: nowrap;
        }}
        td {{
            padding: {'16px 14px' if ppt_export else '12px 10px'};
            border-bottom: 1px solid #e5e7eb;
            text-align: left;
            vertical-align: top;
            line-height: 1.4;
            word-wrap: break-word;
            font-size: 16px;
        }}
        tr:nth-child(even) {{ background: #F9FAFB; }}
        .status-complete {{ color: #16A34A; font-weight: bold; }}
        .status-in-progress {{ color: #F59E0B; font-weight: bold; }}
        .status-not-started {{ color: #6B7280; }}
        .status-on-hold {{ color: #DC2626; font-weight: bold; }}
        
        @media print {{
            .print-toolbar {{ display: none !important; }}
            body {{ padding: 10px; }}
            @page {{ size: landscape; margin: 10mm; }}
            th {{ background: #1E40AF !important; color: white !important; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
            tr:nth-child(even) {{ background: #F9FAFB !important; -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
        }}
    </style>
</head>
<body>
    {'<div class="print-toolbar"><h2>' + table.get("name", "Schedule") + '</h2><button class="btn-print" onclick="window.print()">🖨️ Save as PDF</button><button class="btn-close" onclick="window.close()">✕ Close</button></div>' if is_print_view else ''}
    <h1 class="slide-title">{table.get('name', 'Schedule')}</h1>
    <table>
        <thead>
            <tr>
'''
    
    # Column headers
    for col in columns:
        html += f'                <th>{col.get("header", "Column")}</th>\n'
    
    html += '''            </tr>
        </thead>
        <tbody>
'''
    
    # Data rows
    for row in table.get('rows', []):
        html += '            <tr>\n'
        row_data = row.get('data', {})
        
        for col in columns:
            col_id = col.get('id')
            value = row_data.get(col_id, '')
            
            # Apply status styling if this is a dropdown/status column
            css_class = ''
            if col.get('type') == 'dropdown' and value:
                status_lower = str(value).lower().replace(' ', '-')
                css_class = f'class="status-{status_lower}"'
            
            html += f'                <td {css_class}>{value}</td>\n'
        
        html += '            </tr>\n'
    
    # Empty state
    if not table.get('rows'):
        html += f'            <tr><td colspan="{len(columns)}" style="text-align: center; color: #666; padding: 40px;">No items in this schedule yet.</td></tr>\n'
    
    html += '''        </tbody>
    </table>
'''
    
    # Auto-trigger print dialog for PDF export (not slide view)
    if is_print_view:
        html += '''    <script>
        window.addEventListener('load', function() {
            // Small delay to ensure rendering is complete
            setTimeout(function() { window.print(); }, 400);
        });
    </script>
'''
    
    html += '''</body>
</html>'''
    
    return HTMLResponse(
        content=html,
        status_code=200,
        headers={
            'Cache-Control': 'no-cache, no-store, must-revalidate',
            'Pragma': 'no-cache',
            'Expires': '0'
        }
    )


@router.get("/schedule/print/{project_name}/{table_id}", response_class=HTMLResponse)
async def schedule_print_view(project_name: str, table_id: str):
    """Print-friendly view of a schedule table (export columns only)"""
    return await schedule_table_view(project_name, table_id, export=True, ppt_export=False)


@router.get("/schedule/slide/{project_name}/{table_id}", response_class=HTMLResponse)
async def schedule_slide_view(project_name: str, table_id: str):
    """PowerPoint slide-ready view of a schedule table"""
    return await schedule_table_view(project_name, table_id, export=True, ppt_export=True)


# =============================================================================
# File Import Endpoints
# =============================================================================

def _detect_col_type(header: str, sample_values: list) -> str:
    """Detect whether a column should be typed 'date' or 'text'.

    Returns 'date' if the header name contains date-related keywords OR
    if a majority of sample values look like ISO dates (YYYY-MM-DD).
    Otherwise returns 'text'.
    """
    DATE_KEYWORDS = ('date', 'due', 'start', 'finish', 'end', 'deadline',
                     'target', 'planned', 'by', 'when', 'schedule', 'delivery')
    if any(kw in header.lower() for kw in DATE_KEYWORDS):
        return 'date'
    # Value-heuristic: check first 10 non-empty samples
    iso_pat = re.compile(r'^\d{4}-\d{2}-\d{2}')
    vals = [str(v).strip() for v in sample_values if v is not None and str(v).strip()]
    vals = vals[:10]
    if vals and sum(1 for v in vals if iso_pat.match(v)) > len(vals) // 2:
        return 'date'
    return 'text'


def _normalize_cell_value(cell) -> str:
    """Convert an openpyxl cell value to a clean string.

    datetime/date objects are formatted as YYYY-MM-DD so they are
    recognised as date values by _detect_col_type and the calendar.
    """
    if cell is None:
        return ""
    if isinstance(cell, (_datetime, _date)):
        return cell.strftime('%Y-%m-%d')
    return str(cell).strip()


def parse_excel_file(file_content: bytes) -> tuple:
    """Parse Excel file and return headers and rows"""
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        sheet = workbook.active
        
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return [], []
        
        # First row is headers - safely convert to strings
        headers = []
        for i, h in enumerate(rows[0]):
            if h is None or str(h).strip() == '':
                headers.append(f"Column {i+1}")
            else:
                try:
                    # Convert any type to string safely
                    header_str = str(h).strip()
                    # Python keywords are valid column names
                    headers.append(header_str if header_str else f"Column {i+1}")
                except Exception as e:
                    logger.warning(f"Error converting header at column {i}: {e}")
                    headers.append(f"Column {i+1}")
        
        # Data rows - normalize each cell (datetimes → YYYY-MM-DD, etc.)
        data_rows = []
        for row_idx, row in enumerate(rows[1:], start=2):
            row_data = []
            for col_idx, cell in enumerate(row):
                try:
                    row_data.append(_normalize_cell_value(cell))
                except Exception as e:
                    logger.warning(f"Error converting cell at row {row_idx}, col {col_idx}: {e}")
                    row_data.append("")
            
            # Only add non-empty rows
            if any(str(cell).strip() for cell in row_data):
                data_rows.append(row_data)
        
        logger.info(f"Parsed Excel: {len(headers)} columns, {len(data_rows)} rows")
        return headers, data_rows
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed - cannot parse Excel files")
    except Exception as e:
        logger.error(f"Error parsing Excel file: {e}")
        raise HTTPException(status_code=400, detail=f"Error parsing Excel file: {str(e)}")


def parse_csv_file(file_content: bytes) -> tuple:
    """Parse CSV file and return headers and rows"""
    # Try to decode with different encodings
    for encoding in ['utf-8', 'utf-8-sig', 'latin-1', 'cp1252']:
        try:
            text = file_content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(status_code=400, detail="Could not decode file - unsupported encoding")
    
    # Parse CSV
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    
    if not rows:
        return [], []
    
    headers = rows[0]
    data_rows = rows[1:]
    
    return headers, data_rows


@router.post("/api/schedule/preview")
async def preview_schedule_file(file: UploadFile = File(...)):
    """
    Preview an uploaded file - returns headers and first 5 rows
    """
    try:
        content = await file.read()
        filename = file.filename.lower()
        
        if filename.endswith(('.xlsx', '.xls')):
            headers, data_rows = parse_excel_file(content)
        elif filename.endswith('.csv'):
            headers, data_rows = parse_csv_file(content)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Use Excel (.xlsx) or CSV")
        
        return JSONResponse(content={
            "headers": headers,
            "preview_rows": data_rows[:5],
            "total_rows": len(data_rows)
        })
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error previewing file: {e}")
        raise HTTPException(status_code=400, detail=f"Error reading file: {str(e)}")


@router.post("/api/schedule/{project_name}/import")
async def import_schedule_file(
    project_name: str,
    file: UploadFile = File(...),
    mode: str = Form(...),
    table_name: Optional[str] = Form(None),
    table_id: Optional[str] = Form(None),
    column_mapping: Optional[str] = Form(None)
):
    """
    Import data from an Excel or CSV file into a schedule table
    
    Args:
        project_name: Project to import into
        file: The uploaded file
        mode: 'new' to create a new table, 'existing' to add to existing table
        table_name: Name for new table (if mode=new)
        table_id: ID of existing table (if mode=existing)
        column_mapping: JSON mapping of table column IDs to file column names
    """
    try:
        content = await file.read()
        filename = file.filename.lower()
        
        # Parse file
        if filename.endswith(('.xlsx', '.xls')):
            headers, data_rows = parse_excel_file(content)
        elif filename.endswith('.csv'):
            headers, data_rows = parse_csv_file(content)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format")
        
        if not data_rows:
            raise HTTPException(status_code=400, detail="No data rows found in file")
        
        if mode == 'new':
            # Create new table with auto-detected columns
            if not table_name:
                table_name = file.filename.rsplit('.', 1)[0]
            
            # Create columns from headers - detect dates by name/value heuristics
            columns = []
            for i, header in enumerate(headers):
                sample = [row[i] for row in data_rows[:10] if i < len(row)]
                col_type = _detect_col_type(header, sample)
                columns.append({
                    'header': header,
                    'type': col_type,
                    'width': 150,
                    'visible_in_export': True  # ← ENSURE THIS IS TRUE
                })
            
            # Create the table
            table = schedule_repo.create_table(project_name, table_name, columns)
            table_id = table['id']
            
            # Map headers to column IDs
            header_to_col = {col['header']: col['id'] for col in table['columns']}
            
            # Build all row dicts first, then bulk-insert in a single file write
            bulk_rows = []
            for row_values in data_rows:
                row_data = {}
                for i, value in enumerate(row_values):
                    if i < len(headers):
                        col_id = header_to_col.get(headers[i])
                        if col_id:
                            row_data[col_id] = value
                if any(row_data.values()):  # Skip empty rows
                    bulk_rows.append(row_data)

            rows_imported = schedule_repo.add_rows_bulk(project_name, table_id, bulk_rows)

            return JSONResponse(content={
                "success": True,
                "table_name": table_name,
                "table_id": table_id,
                "rows_imported": rows_imported,
                "columns_created": len(columns)
            })
            
        else:  # mode == 'existing'
            if not table_id:
                raise HTTPException(status_code=400, detail="table_id required for existing mode")
            
            # Get existing table
            table = schedule_repo.get_table(project_name, table_id)
            if not table:
                raise HTTPException(status_code=404, detail="Table not found")
            
            # Parse column mapping
            mapping = {}
            if column_mapping:
                try:
                    mapping = json.loads(column_mapping)
                except:
                    pass
            
            # If no mapping, try to auto-map by header name
            if not mapping:
                for col in table['columns']:
                    for header in headers:
                        if header.lower() == col['header'].lower():
                            mapping[col['id']] = header
                            break
            
            # Reverse mapping: file column -> table column id
            file_to_table = {v: k for k, v in mapping.items()}
            
            # Build all row dicts first, then bulk-insert in a single file write
            bulk_rows = []
            for row_values in data_rows:
                row_data = {}
                for i, value in enumerate(row_values):
                    if i < len(headers):
                        col_id = file_to_table.get(headers[i])
                        if col_id:
                            row_data[col_id] = value
                if any(row_data.values()):  # Skip empty rows
                    bulk_rows.append(row_data)

            rows_imported = schedule_repo.add_rows_bulk(project_name, table_id, bulk_rows)

            return JSONResponse(content={
                "success": True,
                "table_name": table['name'],
                "table_id": table_id,
                "rows_imported": rows_imported
            })
            
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error importing file: {e}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"Import failed due to a server error. Please try again or contact support.")


@router.post("/api/schedule/{project_name}/tables/copy")
async def copy_table_from_program(project_name: str, request: CopyTableRequest):
    """
    Copy a table from another program to the current program.
    
    Args:
        project_name: Target project to copy the table to
        request.source_project: Source project name
        request.source_table_id: ID of the table to copy
        request.new_table_name: Optional new name for the copied table
        request.include_data: Whether to include row data (default: True)
    
    Returns:
        The newly created table
    """
    try:
        # Get source table
        source_table = schedule_repo.get_table(request.source_project, request.source_table_id)
        if not source_table:
            raise HTTPException(status_code=404, detail=f"Source table not found in {request.source_project}")
        
        # Prepare new table data
        new_table_name = request.new_table_name or f"{source_table['name']} (Copy)"
        
        # Create new table with same columns
        new_table = schedule_repo.create_table(project_name, new_table_name, source_table.get('columns', []))
        
        # Copy rows if requested
        rows_copied = 0
        if request.include_data and source_table.get('rows'):
            for row in source_table['rows']:
                # Copy the cell data from the row's 'data' field
                row_data = row.get('data', {})
                schedule_repo.add_row(project_name, new_table['id'], row_data)
                rows_copied += 1
        
        # Get the updated table with rows
        final_table = schedule_repo.get_table(project_name, new_table['id'])
        
        logger.info(f"✅ Copied table '{source_table['name']}' from '{request.source_project}' to '{project_name}' "
                   f"(columns: {len(source_table.get('columns', []))}, rows: {rows_copied})")
        
        return JSONResponse(content={
            "success": True,
            "table": final_table,
            "source_table_name": source_table['name'],
            "rows_copied": rows_copied
        })
        
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error copying table: {e}")
        raise HTTPException(status_code=500, detail=f"Copy failed: {str(e)}")
